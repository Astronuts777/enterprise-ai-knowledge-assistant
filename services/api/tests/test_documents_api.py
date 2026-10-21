from io import BytesIO
from pathlib import Path
from uuid import UUID

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from enterprise_ai_api.main import create_app

ADMIN_A = UUID("00000000-0000-0000-0000-00000000000a")
ADMIN_B = UUID("00000000-0000-0000-0000-00000000000b")
MEMBER_A = UUID("00000000-0000-0000-0000-000000000001")


def headers(user_id: UUID) -> dict[str, str]:
    return {"X-Dev-User-Id": str(user_id)}


def create_workspace(test_client: TestClient, user_id: UUID, name: str) -> tuple[str, str]:
    tenant_response = test_client.post(
        "/v1/tenants", json={"name": name}, headers=headers(user_id)
    )
    assert tenant_response.status_code == 201
    tenant_id = tenant_response.json()["id"]
    space_response = test_client.post(
        f"/v1/tenants/{tenant_id}/spaces",
        json={"name": "Policies", "classification": "confidential"},
        headers=headers(user_id),
    )
    assert space_response.status_code == 201
    return tenant_id, space_response.json()["id"]


def test_upload_preview_publish_and_replace_document_version(tmp_path: Path) -> None:
    app = create_app("sqlite://", storage_root=str(tmp_path / "uploads"))
    with TestClient(app) as test_client:
        tenant_id, space_id = create_workspace(test_client, ADMIN_A, "Acme")
        first_upload = test_client.post(
            f"/v1/tenants/{tenant_id}/spaces/{space_id}/documents",
            data={"title": "Remote work policy"},
            files={
                "file": (
                    "remote-work.md",
                    b"Employees may work remotely.\n\nApproval is required.",
                    "text/markdown",
                )
            },
            headers=headers(ADMIN_A),
        )

        assert first_upload.status_code == 201
        first = first_upload.json()
        assert first["version"]["status"] == "ready"
        assert first["version"]["version_number"] == 1
        document_id = first["document"]["id"]
        version_id = first["version"]["id"]

        chunks = test_client.get(
            f"/v1/tenants/{tenant_id}/document-versions/{version_id}/chunks",
            headers=headers(ADMIN_A),
        )
        assert chunks.status_code == 200
        assert [chunk["locator"] for chunk in chunks.json()] == ["paragraph:1", "paragraph:2"]

        published = test_client.post(
            f"/v1/tenants/{tenant_id}/document-versions/{version_id}/publish",
            headers=headers(ADMIN_A),
        )
        assert published.status_code == 200
        assert published.json()["status"] == "published"

        second_upload = test_client.post(
            f"/v1/tenants/{tenant_id}/documents/{document_id}/versions",
            files={"file": ("remote-work-v2.txt", b"Office approval is no longer required.")},
            headers=headers(ADMIN_A),
        )
        assert second_upload.status_code == 201
        assert second_upload.json()["version_number"] == 2
        second_version_id = second_upload.json()["id"]

        test_client.post(
            f"/v1/tenants/{tenant_id}/document-versions/{second_version_id}/publish",
            headers=headers(ADMIN_A),
        )
        detail = test_client.get(
            f"/v1/tenants/{tenant_id}/documents/{document_id}", headers=headers(ADMIN_A)
        )

    versions = detail.json()["versions"]
    assert [version["status"] for version in versions] == ["published", "superseded"]
    saved_files = [path for path in (tmp_path / "uploads").rglob("*") if path.is_file()]
    assert len(saved_files) == 2


def test_docx_upload_extracts_paragraphs(tmp_path: Path) -> None:
    document = DocxDocument()
    document.add_heading("Security policy", level=1)
    document.add_paragraph("Use multi-factor authentication for administrative access.")
    stream = BytesIO()
    document.save(stream)

    app = create_app("sqlite://", storage_root=str(tmp_path / "uploads"))
    with TestClient(app) as test_client:
        tenant_id, space_id = create_workspace(test_client, ADMIN_A, "Acme")
        uploaded = test_client.post(
            f"/v1/tenants/{tenant_id}/spaces/{space_id}/documents",
            files={
                "file": (
                    "security-policy.docx",
                    stream.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=headers(ADMIN_A),
        )
        version_id = uploaded.json()["version"]["id"]
        chunks = test_client.get(
            f"/v1/tenants/{tenant_id}/document-versions/{version_id}/chunks",
            headers=headers(ADMIN_A),
        )

    assert uploaded.status_code == 201
    assert uploaded.json()["version"]["status"] == "ready"
    assert "multi-factor authentication" in chunks.json()[1]["text"]


def test_failed_parse_is_recorded_and_cannot_be_published(tmp_path: Path) -> None:
    app = create_app("sqlite://", storage_root=str(tmp_path / "uploads"))
    with TestClient(app) as test_client:
        tenant_id, space_id = create_workspace(test_client, ADMIN_A, "Acme")
        uploaded = test_client.post(
            f"/v1/tenants/{tenant_id}/spaces/{space_id}/documents",
            files={"file": ("scan.pdf", b"not a real PDF", "application/pdf")},
            headers=headers(ADMIN_A),
        )
        version = uploaded.json()["version"]
        published = test_client.post(
            f"/v1/tenants/{tenant_id}/document-versions/{version['id']}/publish",
            headers=headers(ADMIN_A),
        )

    assert uploaded.status_code == 201
    assert version["status"] == "failed"
    assert version["error_message"] == "The document could not be parsed"
    assert published.status_code == 409


def test_member_cannot_upload_and_other_tenant_cannot_list(tmp_path: Path) -> None:
    app = create_app("sqlite://", storage_root=str(tmp_path / "uploads"))
    with TestClient(app) as test_client:
        tenant_a, space_a = create_workspace(test_client, ADMIN_A, "Acme")
        tenant_b, _ = create_workspace(test_client, ADMIN_B, "Globex")
        test_client.post(
            f"/v1/tenants/{tenant_a}/members",
            json={"user_id": str(MEMBER_A), "role": "member"},
            headers=headers(ADMIN_A),
        )

        member_upload = test_client.post(
            f"/v1/tenants/{tenant_a}/spaces/{space_a}/documents",
            files={"file": ("policy.txt", b"Policy")},
            headers=headers(MEMBER_A),
        )
        cross_tenant_list = test_client.get(
            f"/v1/tenants/{tenant_a}/spaces/{space_a}/documents",
            headers=headers(ADMIN_B),
        )

    assert tenant_b != tenant_a
    assert member_upload.status_code == 403
    assert cross_tenant_list.status_code == 404
    assert list((tmp_path / "uploads").rglob("*")) == []


def test_local_sqlite_survives_api_restart(tmp_path: Path) -> None:
    database_path = tmp_path / "persistent.db"
    database_url = f"sqlite:///{database_path}"
    storage_root = tmp_path / "uploads"

    first_app = create_app(database_url, storage_root=str(storage_root))
    with TestClient(first_app) as test_client:
        tenant_id, _ = create_workspace(test_client, ADMIN_A, "Persistent Acme")
    first_app.state.engine.dispose()

    second_app = create_app(database_url, storage_root=str(storage_root))
    with TestClient(second_app) as test_client:
        response = test_client.get(f"/v1/tenants/{tenant_id}", headers=headers(ADMIN_A))
    second_app.state.engine.dispose()

    assert database_path.exists()
    assert response.status_code == 200
    assert response.json()["name"] == "Persistent Acme"
