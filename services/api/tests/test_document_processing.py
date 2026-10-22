from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document

from enterprise_ai_api import document_processing as parser
from enterprise_ai_api.retrieval import _best_excerpt


def test_docx_tables_preserve_headers_order_and_row_locators():
    document = Document()
    document.add_paragraph("Expense policy")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(table.rows, [("Item", "Limit"), ("Hotel", "180 AUD"),
                                        ("Meals", "65 AUD")], strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("Approval required before travel.")
    stream = BytesIO()
    document.save(stream)
    chunks = parser.parse_document(stream.getvalue(), ".docx")
    assert [chunk.locator for chunk in chunks] == [
        "paragraph:1", "table:1:row:1", "table:1:row:2", "table:1:row:3", "paragraph:2",
    ]
    assert chunks[2].text == "Item | Limit\nHotel | 180 AUD"
    assert chunks[3].text == "Item | Limit\nMeals | 65 AUD"


def test_mixed_pdf_ocr_does_not_silently_drop_image_pages(monkeypatch):
    pages = [SimpleNamespace(images=[], extract_text=lambda **kw: "Policy summary"),
             SimpleNamespace(images=[object()], extract_text=lambda **kw: "Header only")]
    monkeypatch.setattr(parser, "PdfReader", lambda stream: SimpleNamespace(pages=pages))
    ocr = Mock(return_value="Scanned table: reimbursement limit 180 AUD")
    monkeypatch.setattr(parser, "_ocr_page", ocr)
    chunks = parser.parse_document(b"fixture", ".pdf")
    assert [chunk.locator for chunk in chunks] == ["page:1", "page:2"]
    assert "180 AUD" in chunks[1].text
    assert ocr.call_args.args[:2] == (b"fixture", 2)


def test_pdf_missing_ocr_fails_instead_of_publishing_partial_text(monkeypatch):
    page = SimpleNamespace(images=[object()], extract_text=lambda **kw: "Header only")
    monkeypatch.setattr(parser, "PdfReader", lambda stream: SimpleNamespace(pages=[page]))
    monkeypatch.setattr(parser.shutil, "which", lambda name: None)
    with pytest.raises(parser.DocumentProcessingError, match="未安装 OCR"):
        parser.parse_document(b"fixture", ".pdf")


def test_ocr_cleans_private_temp_files_on_timeout(monkeypatch):
    folders = []
    monkeypatch.setattr(parser.shutil, "which", lambda name: f"/bin/{name}")

    def timeout(command, **kwargs):
        source = Path(command[-2])
        folders.append(source.parent)
        assert source.stat().st_mode & 0o777 == 0o600
        raise parser.subprocess.TimeoutExpired(command, 1)

    monkeypatch.setattr(parser.subprocess, "run", timeout)
    with pytest.raises(parser.DocumentProcessingError, match="OCR 失败或超时"):
        parser._ocr_page(b"private", 1, parser.time.monotonic() + 90)
    assert folders and not folders[0].exists()


def test_excerpt_includes_a_match_near_the_end():
    text = "Unrelated introduction. " * 80 + "The refund deadline is 30 days."
    assert "refund deadline is 30 days" in _best_excerpt(text, {"refund", "deadline"})
